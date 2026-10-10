import zipfile
import re
import sys
from pathlib import Path

ROOT_DIR = Path(__file__).resolve().parents[1]
if hasattr(sys.stdout, "reconfigure"):
    sys.stdout.reconfigure(encoding="utf-8")
docx_path = ROOT_DIR / "documents" / "系統手冊_1.docx"

if not docx_path.exists():
    raise FileNotFoundError(f"找不到檔案: {docx_path}")

print(f"📖 正在讀取並處理: {docx_path}")

# 暫存路徑
temp_docx = docx_path.with_name("系統手冊_1_temp.docx")

replacements_count = 0

with zipfile.ZipFile(docx_path, "r") as zin:
    with zipfile.ZipFile(temp_docx, "w", compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            data = zin.read(item.filename)
            if item.filename == "word/document.xml":
                content = data.decode("utf-8")
                
                # 替換邏輯：匹配 Gemini 相關的 2.5 並改為 3.8
                def replace_func(match):
                    global replacements_count
                    replacements_count += 1
                    return match.group(0).replace("2.5", "3.8")
                
                # 匹配包含 2.5 的 Gemini 字串
                pattern = re.compile(r'(?:Google\s+)?Gemini\s+2\.5(?:\s+Flash)?', re.IGNORECASE)
                new_content = pattern.sub(replace_func, content)
                
                zout.writestr(item, new_content.encode("utf-8"))
            else:
                zout.writestr(item, data)

print(f"✅ 成功替換了 {replacements_count} 處 'Gemini 2.5' 為 'Gemini 3.8'")

# 驗證新產生的 DOCX 能夠被 python-docx 正常解析
import docx
doc = docx.Document(temp_docx)
print(f"📄 驗證成功：Docx 可正常讀取，段落數 {len(doc.paragraphs)}，表格數 {len(doc.tables)}")

# 替換原檔案
temp_docx.replace(docx_path)
print(f"💾 已覆寫更新: {docx_path}")
