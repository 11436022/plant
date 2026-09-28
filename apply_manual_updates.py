import os
import shutil
import zipfile
import docx
from lxml import etree
from PIL import Image

TARGET_DOCX = 'documents/系統手冊_1.docx'
BACKUP_DOCX = 'documents/系統手冊_1_backup.docx'

def step1_backup():
    print("--- STEP 1: Creating Backup ---")
    shutil.copy2(TARGET_DOCX, BACKUP_DOCX)
    print(f"Backed up {TARGET_DOCX} to {BACKUP_DOCX}")

def step2_update_text_and_headings():
    print("--- STEP 2: Updating Headings, TOC, Captions & Text ---")
    doc = docx.Document(TARGET_DOCX)

    # 1. Find and remove empty Heading 1 paragraphs around P126
    for p in list(doc.paragraphs):
        if p.style.name == 'Heading 1' and not p.text.strip():
            print("Removing empty Heading 1 paragraph...")
            p._element.getparent().remove(p._element)

    # Re-fetch paragraphs after removal
    paragraphs = doc.paragraphs

    # 2. Update TOC items for Chapter 1-14
    toc_mapping = {
        '前言\t1': '第1章 前言\t1',
        '營運計畫\t3': '第2章 營運計畫\t3',
        '系統規格\t7': '第3章 系統規格\t7',
        '專案時程與組織分工\t9': '第4章 專案時程與組織分工\t9',
        '需求模型\t15': '第5章 需求模型\t15',
        '設計模型\t28': '第6章 設計模型\t28',
        '實作模型\t30': '第7章 實作模型\t30',
        '資料庫設計\t37': '第8章 資料庫設計\t37',
        '程式\t41': '第9章 程式\t41',
        '測試模型\t42': '第10章 測試模型\t42',
        '操作手冊\t44': '第11章 操作手冊\t44',
        '使用手冊\t44': '第12章 使用手冊\t44',
        '感想\t44': '第13章 感想\t44',
        '參考資料\t44': '第14章 參考資料\t44'
    }

    # Also check if already has chapter numbers or need replacement
    for p in paragraphs[:70]:
        for k, v in toc_mapping.items():
            if p.text.strip() == k.strip():
                p.text = v
                print(f"Updated TOC item: {v}")

    # Fix 表目錄 order if 表 8-2-8 is misplaced before 8-2-6
    # Let's locate the 表目錄 paragraphs
    table_toc_paras = [p for p in paragraphs if '表 8-2-' in p.text]
    # Check if 8-2-8 is present
    p_828 = next((p for p in table_toc_paras if '8-2-8' in p.text), None)
    p_827 = next((p for p in table_toc_paras if '8-2-7' in p.text), None)
    if p_828 and p_827:
        idx_828 = paragraphs.index(p_828)
        idx_827 = paragraphs.index(p_827)
        if idx_828 < idx_827:
            print("Fixing 表 8-2-8 order in 表目錄...")
            # Swap text
            text_828 = p_828.text
            text_827 = p_827.text
            # Wait, let's check all 8-2-5, 8-2-6, 8-2-7, 8-2-8
            p_826 = next((p for p in table_toc_paras if '8-2-6' in p.text), None)
            if p_826:
                # order should be 826, 827, 828
                p_828.text = "表 8-2-6資料表-user_one_time_tokens\t39"
                p_826.text = "表 8-2-7資料表-webcam_alert\t40"
                p_827.text = "表 8-2-8資料表-diagnosis_feedback\t40"
                print("Table of Tables order fixed: 8-2-6, 8-2-7, 8-2-8")

    # 3. Update Heading 1 paragraphs in body
    h1_mapping = [
        ('前言', '第1章 前言'),
        ('營運計畫', '第2章 營運計畫'),
        ('系統規格', '第3章 系統規格'),
        ('專案時程與組織分工', '第4章 專案時程與組織分工'),
        ('需求模型', '第5章 需求模型'),
        ('設計模型', '第6章 設計模型'),
        ('實作模型', '第7章 實作模型'),
        ('資料庫設計', '第8章 資料庫設計'),
        ('程式', '第9章 程式'),
        ('測試模型', '第10章 測試模型'),
        ('操作手冊', '第11章 操作手冊'),
        ('使用手冊', '第12章 使用手冊'),
        ('感想', '第13章 感想'),
        ('參考資料', '第14章 參考資料'),
    ]

    for p in paragraphs:
        if p.style.name == 'Heading 1':
            cleaned = p.text.strip()
            for old_title, new_title in h1_mapping:
                if cleaned == old_title or cleaned.endswith(old_title):
                    p.text = new_title
                    print(f"Updated Heading 1: {new_title}")
                    break

    # 4. Update specific textual descriptions
    for p in paragraphs:
        # P134
        if '備註與結果修正' in p.text:
            p.text = p.text.replace('備註與結果修正', '瀏覽歷史紀錄、編輯診斷備忘筆記與提交反饋糾錯')
            print("Updated P134: 備註與結果修正 -> 編輯診斷備忘筆記與提交反饋糾錯")

        # P136
        if '以 Google Gemini 進行影像初判，再以作物' in p.text:
            p.text = '建立可驗證的雙模型並行 AI 診斷流程：結合 Google Gemini 2.5 Flash 雲端多模態模型與本地自訓特定作物 ConvNet 卷積神經網路（零雲端託管成本），經後端多模型決策仲裁與本地 FAISS 向量知識庫、農業部官方病害蟲害開放資料庫校驗名稱與 70% 信心門檻；處置建議由可追溯資料來源覆核，無法確認時明確標示需人工專家確認。'
            print("Updated P136: Dual-model AI diagnosis pipeline")

        # P142
        if '可串接 Google Gemini，並經資料庫校驗' in p.text:
            p.text = '具備 Google Gemini 2.5 Flash 與本地自訓特定作物 ConvNet 雙模型並行推論、FAISS 向量知識庫校驗、70% 信心門檻及官方資料來源追溯的植物影像診斷流程。'
            print("Updated P142: Dual-model expected result")

        # P170
        if '透過強大的 Google Gemini AI 技術' in p.text:
            p.text = '即時性： 透過 Google Gemini 2.5 Flash 雲端多模態與本地 ConvNet 雙模型並行推論技術，使用者只需透過拍照，即可在短時間內獲得植物健康狀況的精準互補分析，省去繁瑣的網路上網搜尋與比對過程。'
            print("Updated P170: Positioning dual-model")

        # P230
        if 'MySQL 與 Gemini 屬於系統架構元件' in p.text:
            p.text = p.text.replace('MySQL 與 Gemini 屬於系統架構元件', 'MySQL、Gemini 與本地 ConvNet 屬於系統架構元件')
            print("Updated P230: FDD component clarification")

        # P235
        if 'F3 對應 UC-05、UC-06、UC-08（編輯備忘筆記）、UC-09、UC-13（反饋糾錯）' in p.text:
            p.text = p.text.replace(
                'F3 對應 UC-05、UC-06、UC-08（編輯備忘筆記）、UC-09、UC-13（反饋糾錯），F4 對應 UC-10、UC-11，F6 對應 UC-07、UC-09（全站刪除維護）、UC-12、UC-14（回饋審核與標註）',
                'F3 對應 UC-05、UC-06、UC-08（編輯備忘筆記）、UC-13（提交反饋與糾錯），F4 對應 UC-10、UC-11，F6 對應 UC-07、UC-09（管理員刪除維護）、UC-12、UC-14（回饋審核與標註）'
            )
            print("Updated P235: UC-09 belongs exclusively to Admin F6")

        # Caption styles
        if p.text.strip().startswith('表 5-3-13') or p.text.strip().startswith('表 5-3-14') or p.text.strip().startswith('表 8-2-8'):
            try:
                p.style = doc.styles['Caption']
                print(f"Set Caption style on: {p.text.strip()}")
            except Exception as e:
                print(f"Could not set Caption style: {e}")

    # 5. Update Table 19 (表 5-3-9 刪除診斷紀錄)
    t19 = doc.tables[19]
    t19.rows[1].cells[1].text = '系統管理員'
    t19.rows[2].cells[1].text = '系統管理員已完成身分驗證並登入 Web 管理者後台，並位於全站植物診斷紀錄檢閱清單。'
    t19.rows[3].cells[1].text = '指定之異常或測試植物診斷紀錄已由資料庫安全清除，並同步更新全站統計日誌。'
    t19.rows[6].cells[0].text = '管理者於紀錄清單中點擊該筆診斷之「刪除」按鈕，並於二次確認彈窗中點選確定。'
    t19.rows[6].cells[1].text = '前端發送 DELETE API 請求至後端伺服器，自 plant_diary 資料表中安全移除該紀錄，並即時刷新管理列表。'
    print("Updated Table 19 (UC-09): Administrator only")

    doc.save(TARGET_DOCX)
    print("Saved modified text and headings to docx.")

def step3_replace_images_and_fix_extents():
    print("--- STEP 3: Replacing Images & Adjusting Extents in ZIP ---")
    
    # Replacement mapping: image_idx -> png_path
    replacements = {
        12: 'documents/system_activity_diagnosis.png',
        13: 'documents/system_analysis_class.png',
        14: 'documents/system_analysis_object.png',
        15: 'documents/system_sequence_diagnosis.png',
        16: 'documents/system_design_class.png',
        17: 'documents/system_deployment.png',
        18: 'documents/system_package.png',
        19: 'documents/system_component.png',
    }

    # Desired dimensions (EMUs): rId -> (cx, cy)
    # 1 cm = 360000 EMUs
    rid_extents = {
        'rId25': (5400000, 3870000), # image12, ratio 1.395
        'rId26': (5580000, 3302450), # image13, ratio 1.690
        'rId27': (5220000, 3016000), # image14, ratio 1.731
        'rId28': (5580000, 3720000), # image15, ratio 1.50
        'rId29': (5580000, 4141400), # image16, ratio 1.347
        'rId30': (5580000, 3487500), # image17, ratio 1.60
        'rId31': (5580000, 3574690), # image18, ratio 1.561
        'rId32': (5580000, 3534000), # image19, ratio 1.579
        'rId38': (5806440, 3732710), # image25, ratio 1.556
    }

    temp_zip = 'documents/temp_updated.zip'

    with zipfile.ZipFile(TARGET_DOCX, 'r') as zin, zipfile.ZipFile(temp_zip, 'w', compression=zipfile.ZIP_DEFLATED) as zout:
        for item in zin.infolist():
            filename = item.filename
            
            # Check if this file is one of the replaced images
            is_replaced_img = False
            for idx, png_path in replacements.items():
                target_name = f'word/media/image{idx}.png'
                if filename == target_name:
                    print(f"Replacing {filename} with {png_path}...")
                    with open(png_path, 'rb') as f:
                        zout.writestr(item, f.read())
                    is_replaced_img = True
                    break
            
            if is_replaced_img:
                continue

            # Check if this file is document.xml
            if filename == 'word/document.xml':
                print("Updating word/document.xml drawing extents...")
                xml_data = zin.read(filename)
                tree = etree.fromstring(xml_data)
                namespaces = {
                    'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
                    'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
                    'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
                    'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
                    'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships'
                }
                
                drawings = tree.xpath('//w:drawing', namespaces=namespaces)
                for d in drawings:
                    blips = d.xpath('.//a:blip/@r:embed', namespaces=namespaces)
                    if blips and blips[0] in rid_extents:
                        rid = blips[0]
                        new_cx, new_cy = rid_extents[rid]
                        
                        # Update wp:extent
                        for wp_ext in d.xpath('.//wp:extent', namespaces=namespaces):
                            wp_ext.set('cx', str(new_cx))
                            wp_ext.set('cy', str(new_cy))
                            
                        # Update a:xfrm/a:ext
                        for a_ext in d.xpath('.//a:xfrm/a:ext', namespaces=namespaces):
                            a_ext.set('cx', str(new_cx))
                            a_ext.set('cy', str(new_cy))
                            
                        print(f"Updated {rid}: cx={new_cx}, cy={new_cy}")

                updated_xml = etree.tostring(tree, encoding='utf-8', xml_declaration=True)
                zout.writestr(item, updated_xml)
            else:
                zout.writestr(item, zin.read(filename))

    # Replace TARGET_DOCX with temp_zip
    shutil.move(temp_zip, TARGET_DOCX)
    print(f"Successfully repacked {TARGET_DOCX} with new media and updated document.xml!")

if __name__ == '__main__':
    step1_backup()
    step2_update_text_and_headings()
    step3_replace_images_and_fix_extents()
    print("\nALL UPDATES COMPLETED SUCCESSFULLY!")
