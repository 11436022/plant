import sys
import docx

doc = docx.Document('documents/系統手冊_1.docx')
with open('docx_summary.txt', 'w', encoding='utf-8') as f:
    f.write('=== FULL PARAGRAPH DUMP ===\n')
    for i, p in enumerate(doc.paragraphs):
        text = p.text.strip()
        if text:
            f.write(f'P{i:03d} [{p.style.name}]: {text}\n')
    f.write('\n=== TABLES DUMP ===\n')
    for t_idx, table in enumerate(doc.tables):
        rows = len(table.rows)
        cols = len(table.columns)
        header = [cell.text.strip().replace('\n', ' ') for cell in table.rows[0].cells]
        f.write(f"Table {t_idx} ({rows}x{cols}): {' | '.join(header[:5])}\n")
print('Done dumping')
