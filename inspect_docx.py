import sys
import docx
import re

sys.stdout.reconfigure(encoding='utf-8')
doc = docx.Document('documents/系統手冊_1.docx')

print(f"Total paragraphs: {len(doc.paragraphs)}")
print(f"Total tables: {len(doc.tables)}")

rels = doc.part.rels
img_map = {}
for rel_id, rel in rels.items():
    if "image" in rel.target_ref:
        img_map[rel_id] = rel.target_ref

print("\n--- IMAGES IN DOCUMENT ---")
for i, p in enumerate(doc.paragraphs):
    xml = p._p.xml
    if 'blip' in xml:
        embeds = re.findall(r'r:embed="([^"]+)"', xml)
        prev_p = doc.paragraphs[i-1].text if i > 0 else ''
        next_p = doc.paragraphs[i+1].text if i < len(doc.paragraphs)-1 else ''
        targets = [img_map.get(e, e) for e in embeds]
        print(f"P{i}: targets={targets}")
        print(f"   Prev text: {prev_p[:80]}")
        print(f"   Curr text: {p.text[:80]}")
        print(f"   Next text: {next_p[:80]}")
