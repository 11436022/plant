import docx
import zipfile
from PIL import Image
import io

doc_path = 'documents/系統手冊_1.docx'
doc = docx.Document(doc_path)

print("=== 1. VERIFYING HEADINGS IN BODY ===")
h1s = [p.text for p in doc.paragraphs if p.style.name == 'Heading 1']
for i, h in enumerate(h1s, 1):
    print(f"  {i:02d}: {h}")

print("\n=== 2. VERIFYING TOC ENTRIES ===")
for i, p in enumerate(doc.paragraphs[16:62]):
    if '章' in p.text:
        print(f"  P{i+16}: {p.text}")

print("\n=== 3. VERIFYING ALL CHAPTER 7 IMAGES ===")
with zipfile.ZipFile(doc_path, 'r') as z:
    for i in range(17, 25):
        img_name = f'word/media/image{i}.png'
        data = z.read(img_name)
        img = Image.open(io.BytesIO(data))
        print(f"  image{i}.png: size={img.size}, bytes={len(data)}")

print("\nALL VERIFICATIONS COMPLETE!")
