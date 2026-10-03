import docx
import shutil
import os
import zipfile

docx_path = 'documents/系統手冊_1.docx'
backup_path = 'documents/系統手冊_1.docx.bak_plan_a'

print("Creating backup...")
shutil.copyfile(docx_path, backup_path)

doc = docx.Document(docx_path)

# 1. Update Table 5-1-1 (FR-04)
# Let's find FR-04 in tables
updated_fr04 = False
for t in doc.tables:
    for row in t.rows:
        for i, cell in enumerate(row.cells):
            if 'FR-04' in cell.text:
                desc_cell = row.cells[1] if len(row.cells) > 1 else cell
                print(f"Found FR-04: {desc_cell.text}")
                desc_cell.text = "使用者確認診斷後，系統才保存圖片、結果與備註；若診斷為未知作物或無法判定病害，系統嚴格禁止存入病歷日記，僅保留結果有誤（回饋）選項。"
                updated_fr04 = True
                break
        if updated_fr04:
            break
    if updated_fr04:
        break

# 2. Update Table 5-3-4 (上傳圖片並取得分析結果)
updated_t534 = False
for t in doc.tables:
    table_text = "".join(c.text for row in t.rows for c in row.cells)
    if '上傳圖片並取得分析結果' in table_text and '結束狀態' in table_text:
        for row in t.rows:
            row_str = " | ".join(c.text for c in row.cells)
            if '(等待結束)' in row_str:
                for c in row.cells:
                    if 'App 顯示' in c.text:
                        print(f"Found Table 5-3-4 response: {c.text}")
                        c.text = "App 顯示校驗後的結果、信心度、建議與來源；若為未知作物或無法判定，App 自動隱藏儲存按鈕，僅保留「結果有誤？（回饋）」；若為已知確診病害，則同時提供儲存與回饋選項。"
                        updated_t534 = True
                        break
            if '結束狀態' in row_str:
                for c in row.cells:
                    if '使用者看到經資料庫校驗' in c.text:
                        print(f"Found Table 5-3-4 end state: {c.text}")
                        c.text = "使用者看到經資料庫校驗的診斷結果與可追溯來源；若為未知/無法判定則僅提供糾錯反饋選項。"
                        break
        if updated_t534:
            break

# 3. Update Table 5-3-5 (儲存診斷紀錄)
updated_t535 = False
for t in doc.tables:
    table_text = "".join(c.text for row in t.rows for c in row.cells)
    if '儲存診斷紀錄' in table_text and '前提' in table_text and '使用者筆記' in table_text:
        for row in t.rows:
            row_str = " | ".join(c.text for c in row.cells)
            if '前提' in row_str:
                for c in row.cells:
                    if '診斷分析與資料庫校驗已完成' in c.text:
                        print(f"Found Table 5-3-5 precondition: {c.text}")
                        c.text = "診斷分析已完成，且診斷結果為已知確診病害（非未知作物或無法判定之結果）。"
                        updated_t535 = True
                        break
            if '點擊「確認儲存」按鈕' in row_str:
                for c in row.cells:
                    if 'App 將所有資訊打包傳送至後端' in c.text:
                        c.text = "App 將所有資訊打包傳送至後端；若為未知病害則後端阻擋並回傳 HTTP 400 錯誤。"
                        break
        if updated_t535:
            break

# 4. Update Table 5-3-13 (提交診斷反饋與糾錯)
updated_t5313 = False
for t in doc.tables:
    table_text = "".join(c.text for row in t.rows for c in row.cells)
    if '提交診斷反饋與糾錯' in table_text and '前提' in table_text:
        for row in t.rows:
            row_str = " | ".join(c.text for c in row.cells)
            if '前提' in row_str:
                for c in row.cells:
                    if '使用者已完成植物影像診斷' in c.text:
                        print(f"Found Table 5-3-13 precondition: {c.text}")
                        c.text = "使用者完成診斷，包含兩類情境：(1) 診斷結果為未知或無法判定，系統僅提供反饋糾錯選項；(2) 診斷雖有確診結果，但使用者認為推論有誤。"
                        updated_t5313 = True
                        break
        if updated_t5313:
            break

print(f"Table updates: FR-04={updated_fr04}, T5-3-4={updated_t534}, T5-3-5={updated_t535}, T5-3-13={updated_t5313}")
doc.save(docx_path)
print("Saved modified text into docx!")

# 5. Replace the 3 images: image12.png, image15.png, image22.png
image_replacements = {
    'word/media/image12.png': 'documents/system_activity_diagnosis.png',
    'word/media/image15.png': 'documents/system_sequence_diagnosis.png',
    'word/media/image22.png': 'documents/system_state_save_diary.png'
}

temp_docx = 'documents/系統手冊_1_temp.docx'
zin = zipfile.ZipFile(docx_path, 'r')
zout = zipfile.ZipFile(temp_docx, 'w', compression=zipfile.ZIP_DEFLATED)

replaced_count = 0
for item in zin.infolist():
    buffer = zin.read(item.filename)
    if item.filename in image_replacements:
        src_png = image_replacements[item.filename]
        with open(src_png, 'rb') as f:
            new_bytes = f.read()
        print(f"Replacing {item.filename} with {src_png} ({len(buffer)} -> {len(new_bytes)} bytes)...")
        zout.writestr(item, new_bytes)
        replaced_count += 1
    else:
        zout.writestr(item, buffer)

zin.close()
zout.close()

if replaced_count == 3:
    os.replace(temp_docx, docx_path)
    if os.path.exists(backup_path):
        os.remove(backup_path)
    print("Successfully replaced all 3 images in documents/系統手冊_1.docx!")
else:
    print(f"Warning: Only replaced {replaced_count}/3 images!")
    if os.path.exists(temp_docx):
        os.remove(temp_docx)
