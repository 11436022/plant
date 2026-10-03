import docx
import zipfile
import sys
import re

doc = docx.Document('documents/系統手冊_1.docx')
with zipfile.ZipFile('documents/系統手冊_1.docx', 'r') as z:
    rels = z.read('word/_rels/document.xml.rels').decode('utf-8')

r_map = {}
for line in rels.split('<Relationship '):
    if 'Target="media/' in line:
        rid = line.split('Id="')[1].split('"')[0]
        target = line.split('Target="')[1].split('"')[0]
        r_map[rid] = target

print("Total images in rels:", len(r_map))

for i, p in enumerate(doc.paragraphs):
    text = p.text.strip()
    if re.match(r'^圖\s*\d+-\d+-\d+', text):
        # caption paragraph
        prev_xml = doc.paragraphs[i-1]._element.xml if i > 0 else ''
        curr_xml = p._element.xml
        next_xml = doc.paragraphs[i+1]._element.xml if i+1 < len(doc.paragraphs) else ''
        
        found = []
        for rid, target in r_map.items():
            if f'r:embed="{rid}"' in prev_xml or f'r:embed="{rid}"' in curr_xml or f'r:embed="{rid}"' in next_xml:
                found.append((rid, target))
        print(f"{text} --> {found}")
